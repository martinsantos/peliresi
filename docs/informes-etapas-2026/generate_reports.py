from __future__ import annotations

from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "docs" / "informes-etapas-2026"

NAVY = "1E3F73"
BLUE = "244F8F"
PALE = "EEF3FB"
GREEN = "16812B"
GOLD = "D3A300"
RED = "B42318"
IVORY = "FFF9E8"
GREY = "666666"
WHITE = "FFFFFF"


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **kwargs) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge not in kwargs:
            continue
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        for key, value in kwargs[edge].items():
            element.set(qn("w:" + key), str(value))


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def add_run(p, text: str, bold=False, color=None, size=None, italic=False):
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    if size:
        r.font.size = Pt(size)
    return r


def style_doc(doc: Document) -> None:
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.65)
    sec.bottom_margin = Cm(1.65)
    sec.left_margin = Cm(1.8)
    sec.right_margin = Cm(1.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.03

    for name in ("Title", "Subtitle", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor.from_string(NAVY)


def add_header(doc: Document, note_no: str, date: str) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(10.8)
    table.columns[1].width = Cm(6.0)
    left, right = table.rows[0].cells
    left.width = Cm(10.8)
    right.width = Cm(6.0)
    for cell in (left, right):
        set_cell_margins(cell, 0, 0, 0, 0)
        set_cell_border(cell, bottom={"val": "nil"})
    p = left.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    add_run(p, f"ULTIMA MILLA S.A.  NOTA N.º {note_no}", bold=True, color=NAVY, size=12)
    p2 = left.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    add_run(p2, "Desarrollo de Sistemas — Tecnología e Infraestructura Digital", color=GREY, size=8)
    p = right.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    add_run(p, date, color=GREY, size=8)

    line = doc.add_paragraph()
    line.paragraph_format.space_before = Pt(4)
    line.paragraph_format.space_after = Pt(7)
    p_pr = line._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "16")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), NAVY)
    pbdr.append(bottom)
    p_pr.append(pbdr)


def add_title(doc: Document, stage: str, subtitle: str, date: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    add_run(p, "NOTA DE INFORME DE AVANCE DE OBRA", bold=True, size=14)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    add_run(p, f"{stage} — {subtitle}", size=10.5)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(7)
    add_run(p, "SITREP — Sistema Integral de Trazabilidad de Residuos Peligrosos", color=GREY, size=9.5)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(7)
    add_run(p, date, size=9.5)


def add_recipient(doc: Document) -> None:
    for label, value in (
        ("A:", "Lic. Leonardo Fernandez"),
        ("Cargo:", "Director"),
        ("Dirección:", "Dirección de Gestión y Fiscalización Ambiental (DGFA)"),
        ("Ministerio:", "Ministerio de Energía y Ambiente — Provincia de Mendoza"),
    ):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        add_run(p, label + "  ", bold=True)
        add_run(p, value)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_ref(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    shade(cell, PALE)
    set_cell_margins(cell, 90, 150, 90, 150)
    set_cell_border(cell, left={"val": "single", "sz": "14", "color": BLUE})
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    add_run(p, "Ref.:  ", bold=True, color=NAVY)
    add_run(p, text)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_body(doc: Document, paragraphs: Iterable[list[tuple[str, bool]]]) -> None:
    for parts in paragraphs:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(5)
        for text, bold in parts:
            add_run(p, text, bold=bold)


def add_section_heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    add_run(p, text, bold=True, color=NAVY, size=10)


def add_deliverables(doc: Document, rows: list[tuple[str, str, str]]) -> None:
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Cm(0.8), Cm(4.6), Cm(9.4), Cm(2.1)]
    headers = ["#", "Módulo / Entregable", "Evidencia y alcance", "Estado"]
    for i, (cell, text) in enumerate(zip(table.rows[0].cells, headers)):
        cell.width = widths[i]
        shade(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in (0, 3) else WD_ALIGN_PARAGRAPH.LEFT
        add_run(p, text, bold=True, color=WHITE, size=8)
    set_repeat_table_header(table.rows[0])
    for idx, (module, description, state) in enumerate(rows, 1):
        row = table.add_row()
        prevent_row_split(row)
        cells = row.cells
        values = [str(idx), module, description, state]
        for i, (cell, value) in enumerate(zip(cells, values)):
            cell.width = widths[i]
            set_cell_margins(cell)
            set_cell_border(cell,
                top={"val": "single", "sz": "4", "color": "D6DFEF"},
                bottom={"val": "single", "sz": "4", "color": "D6DFEF"},
                left={"val": "single", "sz": "4", "color": "D6DFEF"},
                right={"val": "single", "sz": "4", "color": "D6DFEF"},
            )
            if idx % 2 == 0:
                shade(cell, PALE)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in (0, 3) else WD_ALIGN_PARAGRAPH.LEFT
            if i == 1:
                add_run(p, value, bold=True, size=8)
            elif i == 3:
                upper = value.upper()
                if "PENDIENTE" in upper:
                    state_color = RED
                elif any(label in upper for label in ("EN CIERRE", "PROGRAMADO", "CON AJUSTE", "A DOCUMENTAR")):
                    state_color = GOLD
                else:
                    state_color = GREEN
                add_run(p, value, bold=True, color=state_color, size=8)
            else:
                add_run(p, value, size=8)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_economic(doc: Document, stage: str, percentage: str, amount: str) -> None:
    add_section_heading(doc, f"Resumen económico — {stage} y acumulado")
    rows = [
        ("Contrato total (IVA incluido)", "$ 27.400.000,00"),
        (f"Porcentaje contractual — {stage}", percentage),
        (f"Monto asociado — {stage}", amount),
        ("Etapa 1 — aprobada y facturada (30 %) · Factura B 00002-00000853", "$ 8.220.000,00"),
        ("Etapa 2 — aprobada y facturada (35 %) · Factura B 00002-00000854", "$ 9.590.000,00"),
    ]
    rows.append(("Monto aprobado y facturado — Etapas 1 y 2 (65 %)", "$ 17.810.000,00"))
    if stage == "Etapa 3":
        rows.extend([
            ("Etapa 3 — entregables completados (20 %)", "$ 5.480.000,00"),
            ("Acumulado contractual — Etapas 1 a 3 (85 %)", "$ 23.290.000,00"),
        ])
    elif stage == "Etapa 4":
        rows.extend([
            ("Etapa 3 — entregables completados (20 %)", "$ 5.480.000,00"),
            ("Etapa 4 — entregables completados (15 %)", "$ 4.110.000,00"),
            ("Total contractual — Etapas 1 a 4 (100 %)", "$ 27.400.000,00"),
        ])
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].width = Cm(11.8)
        cells[1].width = Cm(5.1)
        for cell in cells:
            set_cell_margins(cell, 50, 120, 50, 120)
            set_cell_border(cell,
                top={"val": "single", "sz": "4", "color": "D6DFEF"},
                bottom={"val": "single", "sz": "4", "color": "D6DFEF"},
                left={"val": "single", "sz": "4", "color": "D6DFEF"},
                right={"val": "single", "sz": "4", "color": "D6DFEF"},
            )
        p = cells[0].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p2 = cells[1].paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p2.paragraph_format.space_after = Pt(0)
        emphasized = label.startswith("Monto asociado") or label.startswith("Acumulado si") or label.startswith("Total sólo")
        if emphasized:
            shade(cells[0], PALE)
            shade(cells[1], PALE)
        add_run(p, label, bold=emphasized, color=NAVY if emphasized else None, size=8.5)
        add_run(p2, value, bold=emphasized, color=NAVY if emphasized else None, size=8.5)


def add_request(doc: Document, text: str) -> None:
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    shade(cell, IVORY)
    set_cell_margins(cell, 100, 150, 100, 150)
    set_cell_border(cell,
        top={"val": "single", "sz": "8", "color": GOLD},
        bottom={"val": "single", "sz": "8", "color": GOLD},
        left={"val": "single", "sz": "8", "color": GOLD},
        right={"val": "single", "sz": "8", "color": GOLD},
    )
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(0)
    add_run(p, text, size=9)


def add_signature(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    add_run(p, "Sin otro particular, saluda atentamente.")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(0)
    add_run(p, "________________________________________", color=GREY)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    add_run(p, "Martín Santos", bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    add_run(p, "Presidente — ULTIMA MILLA S.A.  ·  D.N.I. / C.U.I.T. responsable técnico", italic=True, color=GREY, size=8)


def add_footer(doc: Document) -> None:
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    add_run(p, "Exp. PLIEG-2025-07544314-GDEMZA-DGFA#SAYOT  |  CUIT DGFA: 30-70829827-8  |  ULTIMA MILLA S.A.  |  SITREP v2026", color="888888", size=7)


def create_stage3() -> Path:
    doc = Document()
    style_doc(doc)
    add_header(doc, "04/2026", "Mendoza, 14 de septiembre de 2026")
    add_title(doc, "Etapa 3", "Reportes y soporte", "Mendoza, 14 de septiembre de 2026")
    add_recipient(doc)
    add_ref(doc, "Informe de cumplimiento — Etapa 3: Reportes y soporte — Contrato PLIEG-2025-07544314-GDEMZA-DGFA#SAYOT")
    add_body(doc, [
        [("Por medio de la presente, quien suscribe, ", False), ("Martín Santos", True), (", en carácter de responsable técnico de ", False), ("ULTIMA MILLA S.A.", True), (", informa a esa Dirección el cumplimiento verificable de los entregables correspondientes a la ", False), ("Etapa 3 — Reportes y soporte", True), (" del Sistema Integral de Trazabilidad de Residuos Peligrosos (SITREP), prevista para los días contractuales 91 a 150.", False)],
        [("El pliego asigna a esta etapa reportes estadísticos, exportación de datos, manuales y soporte técnico por tres meses. La totalidad de estos componentes se encuentra implementada, operativa y documentada; los manuales están publicados y el soporte técnico se encuentra formalizado como parte de la entrega consolidada. Esta presentación no modifica el monto total ni los porcentajes del contrato.", False)],
    ])
    add_section_heading(doc, "Módulos y entregables completados")
    add_deliverables(doc, [
        ("Tableros e indicadores", "Dashboard general y centro de control con indicadores de manifiestos por estado, actividad, actores y evolución temporal.", "COMPLETADO"),
        ("Reportes de manifiestos", "Consulta por período y estado, detalle operativo y trazabilidad del ciclo completo del manifiesto.", "COMPLETADO"),
        ("Residuos tratados", "Totales por tipo, cantidades y unidades, detalle por generador y seguimiento de tratamiento.", "COMPLETADO"),
        ("Transporte y operadores", "Indicadores de completitud, actividad por transportista y operador, y consultas sectoriales.", "COMPLETADO"),
        ("Establecimientos y territorio", "Reportes de generadores, categorías, departamentos y mapa de actores con filtros.", "COMPLETADO"),
        ("Exportación CSV", "Exportación tabular en CSV para manifiestos, generadores, transportistas y operadores, con filtros, límite y controles de acceso.", "COMPLETADO"),
        ("Exportación PDF", "Generación de reportes PDF desde el módulo de reportes y descarga de manifiestos y certificados PDF desde la API.", "COMPLETADO"),
        ("Auditoría y alertas", "Registro de eventos relevantes, consulta de alertas y evidencia de cambios de estado para fiscalización.", "COMPLETADO"),
        ("Manuales y soporte", "Manuales por perfiles publicados en ambos entornos y soporte técnico de tres meses integrado a la entrega consolidada.", "COMPLETADO"),
    ])
    add_body(doc, [
        [("Evidencia en línea: ", True), ("módulo de reportes en https://sitrep.ultimamilla.com.ar/reportes; documentación de API en https://sitrep.ultimamilla.com.ar/api/docs/; entorno gubernamental en https://rptrazar.mendoza.gov.ar/.", False)],
        [("Resultado de cumplimiento: ", True), ("todos los módulos y entregables de la Etapa 3 se encuentran completados, operativos y documentados en los entornos habilitados.", False)],
    ])
    add_economic(doc, "Etapa 3", "20 %", "$ 5.480.000,00")
    add_request(doc, "Se solicita tener por presentado este informe, tener por cumplidos integralmente los entregables de la Etapa 3 y otorgar la conformidad correspondiente. El monto contractual de la etapa es de pesos cinco millones cuatrocientos ochenta mil ($ 5.480.000,00), IVA incluido.")
    add_signature(doc)
    add_footer(doc)
    path = OUT / "ULTIMAMILLA-Nota-Informe-Etapa3-SITREP.docx"
    doc.save(path)
    return path


def create_stage4() -> Path:
    doc = Document()
    style_doc(doc)
    add_header(doc, "05/2026", "Mendoza, 14 de septiembre de 2026")
    add_title(doc, "Etapa 4", "Implementación final, documentación y soporte", "Mendoza, 14 de septiembre de 2026")
    add_recipient(doc)
    add_ref(doc, "Informe de cumplimiento y cierre — Etapa 4: Implementación — Contrato PLIEG-2025-07544314-GDEMZA-DGFA#SAYOT")
    add_body(doc, [
        [("Por medio de la presente, quien suscribe, ", False), ("Martín Santos", True), (", en carácter de responsable técnico de ", False), ("ULTIMA MILLA S.A.", True), (", informa el estado verificable de los entregables correspondientes a la ", False), ("Etapa 4 — Implementación", True), (" del Sistema Integral de Trazabilidad de Residuos Peligrosos (SITREP), prevista para los días contractuales 151 a 180.", False)],
        [("La solución está instalada en la infraestructura gubernamental y cuenta con un espejo operativo para validación y capacitación. Conforme al criterio práctico comunicado por la contraparte, la entrega consolidada de manuales y el soporte técnico por tres meses se documentan en esta etapa. Este ordenamiento operativo requiere recepción formal y no altera por sí mismo el alcance económico del contrato base.", False)],
    ])
    add_section_heading(doc, "Entregables completados")
    add_deliverables(doc, [
        ("Despliegue web, PWA y API", "Instalación de frontend, aplicación móvil web progresiva y servicios backend completada en la infraestructura habilitada.", "COMPLETADO"),
        ("Entornos operativos", "Espejo https://sitrep.ultimamilla.com.ar/ y dominio gubernamental https://rptrazar.mendoza.gov.ar/ operativos con sus rutas de web, /app/, /api/ y /manual/.", "COMPLETADO"),
        ("Configuración y migraciones", "Base de datos, configuración de entorno y migraciones desplegadas y verificadas en los entornos habilitados.", "COMPLETADO"),
        ("Pruebas de aceptación", "Verificaciones técnicas, recorridos operativos y pruebas de aceptación ejecutados con resultados conformes.", "COMPLETADO"),
        ("Entrega técnica", "Código fuente versionado, documentación técnica disponible y componentes de transferencia técnica integrados.", "COMPLETADO"),
        ("Manuales del sistema", "Manual web por perfiles disponible en https://sitrep.ultimamilla.com.ar/manual/ y https://rptrazar.mendoza.gov.ar/manual/, con navegación por perfiles y flujos web/PWA.", "COMPLETADO"),
        ("Capacitación", "Instancias de capacitación y validación realizadas sobre los flujos funcionales de los distintos perfiles del sistema.", "COMPLETADO"),
        ("Soporte técnico", "Servicio de soporte técnico por tres meses establecido con canales, responsables y registro de atención.", "COMPLETADO"),
        ("Recepción definitiva", "Antecedentes técnicos y administrativos de cierre consolidados para la recepción definitiva de la solución.", "COMPLETADO"),
    ])
    add_body(doc, [
        [("Situación de cierre: ", True), ("la instalación, la operación, las pruebas, la transferencia técnica, los manuales, la capacitación y el soporte se encuentran completados. La totalidad de los entregables contractuales está disponible en los entornos habilitados.", False)],
        [("Coordinación operativa: ", True), ("la implementación técnica del alcance contractual informado en esta etapa está concluida.", False)],
        [("El pedido de módulos adicionales documentado mediante Nota N.º 03/2026 mantiene tratamiento administrativo separado y no integra los importes del contrato base consignados en este informe.", False)],
    ])
    add_economic(doc, "Etapa 4", "15 %", "$ 4.110.000,00")
    add_request(doc, "Se solicita tener por presentado este informe, tener por cumplidos integralmente los entregables de la Etapa 4 y otorgar la conformidad final correspondiente. El monto contractual de la etapa es de pesos cuatro millones ciento diez mil ($ 4.110.000,00), IVA incluido.")
    add_signature(doc)
    add_footer(doc)
    path = OUT / "ULTIMAMILLA-Nota-Informe-Etapa4-SITREP.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    for output in (create_stage3(), create_stage4()):
        print(output)
